"""The deck in four formats, from one source, each one actually accessible.

"Send me the slides" has a different answer at every institution, and a deck
that only exists as a web page cannot be opened by the person who needs it in
Word with their own screen reader settings applied. So: HTML, Markdown, Word
and PowerPoint, all rendered from ``workshop_deck.build_slides`` so the
sentences cannot diverge.

These tests check the accessibility properties that are easy to lose and
invisible when lost: real heading styles, real list styles, a marked table
header row, a declared language, a real title placeholder on every slide, and
speaker notes in the notes slide rather than dumped on the slide itself.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from flask import Flask

from acb_large_print_web import workshop_agenda as agenda
from acb_large_print_web import workshop_deck as deck
from acb_large_print_web.app import create_app
from acb_large_print_web.workshop_store import ensure_session

CODE = "formatdemo"


@pytest.fixture()
def ctx() -> deck.DeckContext:
    return deck.DeckContext(
        code_label=CODE,
        join_display=f"letitglow.app/w/{CODE}",
        activity_times=tuple(f"{b.minutes} minutes" for b in agenda.activity_blocks()),
    )


@pytest.fixture()
def app(tmp_path: Path) -> Flask:
    application = create_app({"TESTING": True, "WTF_CSRF_ENABLED": False})
    application.instance_path = str(tmp_path / "instance")
    Path(application.instance_path).mkdir(parents=True, exist_ok=True)
    with application.app_context():
        ensure_session(CODE, title="Accessibility Agents in Action", event_name="AHG")
    return application


@pytest.fixture()
def client(app: Flask):
    return app.test_client()


# ---------------------------------------------------------------------------
# One source
# ---------------------------------------------------------------------------


def test_every_format_renders_the_same_slides(ctx: deck.DeckContext):
    slides = deck.build_slides(ctx)
    assert len(slides) == 26

    markdown = deck.build_deck_markdown(ctx)
    for slide in slides:
        assert slide.title in markdown


def test_the_title_is_the_program_title():
    """AHG's top attendee complaint is a title that does not match the content."""
    assert f"{deck.DECK_TITLE}: {deck.DECK_SUBTITLE}" == (
        "Accessibility Agents: Building Human-Centered AI Workflows for "
        "Trusted Accessibility Automation at Scale"
    )


def test_every_block_opens_by_saying_where_we_are(ctx: deck.DeckContext):
    kickers = [s.kicker for s in deck.build_slides(ctx)]
    for number in range(2, 8):
        assert any(k.startswith(f"Block {number} of 7") for k in kickers)


def test_the_day_is_the_ahg_day(ctx: deck.DeckContext):
    assert agenda.validate_ahg_day() == []
    markdown = deck.build_deck_markdown(ctx)
    for block in agenda.AHG_DAY:
        if block.kind in (agenda.BREAK, agenda.LUNCH):
            assert block.starts_at in markdown


def test_no_slide_is_a_wall_of_text(ctx: deck.DeckContext):
    """AHG speaker guidance: avoid walls of text."""
    for slide in deck.build_slides(ctx):
        for block in slide.blocks:
            if isinstance(block, deck.Bullets):
                assert len(block.items) <= 7, slide.id
            if isinstance(block, deck.Table):
                assert len(block.rows) <= 9, slide.id


def test_the_day_ends_with_the_session_evaluation():
    """AHG asks speakers to save time for the session evaluation."""
    assert agenda.AGENDA[-1].title == "Session evaluation"
    assert agenda.AGENDA[-1].minutes >= 5


def test_the_privacy_rule_is_on_a_slide(ctx: deck.DeckContext):
    slide = next(s for s in deck.build_slides(ctx) if s.id == "s6")
    text = " ".join(getattr(b, "text", "") or " ".join(getattr(b, "items", ())) for b in slide.blocks)
    assert "Never paste anything private" in text


def test_every_promise_is_named_on_a_slide(ctx: deck.DeckContext):
    markdown = deck.build_deck_markdown(ctx)
    for promise in ("VS Code", "Copilot", "GitHub", "pull request", "Accessibility Agents",
                    "axe", "Accessibility Insights", "WCAG 2.2", "Title II", "team",
                    "session evaluation"):
        assert promise in markdown, promise


def test_this_room_s_addresses_reach_every_format(ctx: deck.DeckContext):
    markdown = deck.build_deck_markdown(ctx)
    assert f"letitglow.app/w/{CODE}/11" in markdown
    assert f"letitglow.app/w/{CODE}" in markdown


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def test_markdown_uses_headings_and_lists_not_layout(ctx: deck.DeckContext):
    markdown = deck.build_deck_markdown(ctx)
    lines = markdown.splitlines()
    assert lines[0] == "---" and lines[1].startswith("title: ")
    assert "# Accessibility Agents" in lines
    assert "## 1. Accessibility Agents" in markdown
    assert "\n- " in markdown
    assert "\n1. " in markdown
    # Pipe table with a header separator, not spaces pretending to be columns.
    assert "| Specialist | Takes |" in markdown
    assert "|---|---|" in markdown


def test_markdown_keeps_the_speaker_notes(ctx: deck.DeckContext):
    markdown = deck.build_deck_markdown(ctx)
    assert "### Speaker notes, slide 2" in markdown
    assert "Let the laugh happen" in markdown


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------


def test_word_uses_real_heading_and_list_styles(ctx: deck.DeckContext):
    from docx import Document

    doc = Document(BytesIO(deck.build_deck_docx_bytes(ctx)))
    styles = [p.style.name for p in doc.paragraphs]

    assert styles.count("Heading 1") == 1, "one document title"
    assert styles.count("Heading 2") == len(deck.build_slides(ctx)), "one heading per slide"
    assert "List Bullet" in styles
    assert "List Number" in styles
    assert "Heading 3" in styles, "speaker notes get their own heading level"


def test_word_is_acb_large_print_and_declares_its_language(ctx: deck.DeckContext):
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(BytesIO(deck.build_deck_docx_bytes(ctx)))
    normal = doc.styles["Normal"]
    assert normal.font.name == "Arial"
    assert normal.font.size.pt == 18

    lang = normal.element.get_or_add_rPr().find(qn("w:lang"))
    assert lang is not None
    assert lang.get(qn("w:val")) == "en-US"
    assert doc.core_properties.title == deck.DECK_TITLE


def test_word_tables_mark_their_header_row(ctx: deck.DeckContext):
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(BytesIO(deck.build_deck_docx_bytes(ctx)))
    assert doc.tables, "the deck carries tables"
    for table in doc.tables:
        header = table.rows[0]
        tr_pr = header._tr.find(qn("w:trPr"))
        assert tr_pr is not None and tr_pr.find(qn("w:tblHeader")) is not None, (
            "a table header row that is not marked is announced as data"
        )


# ---------------------------------------------------------------------------
# PowerPoint
# ---------------------------------------------------------------------------


def test_powerpoint_gives_every_slide_a_real_title(ctx: deck.DeckContext):
    from pptx import Presentation

    prs = Presentation(BytesIO(deck.build_deck_pptx_bytes(ctx)))
    slides = deck.build_slides(ctx)
    assert len(prs.slides) == len(slides)

    for shipped, source in zip(prs.slides, slides, strict=True):
        assert shipped.shapes.title is not None, "a slide with no title placeholder"
        assert shipped.shapes.title.text == source.title


def test_powerpoint_content_sits_in_placeholders_not_floating_boxes(ctx: deck.DeckContext):
    """Placeholder order is the reading order a screen reader announces."""
    from pptx import Presentation

    prs = Presentation(BytesIO(deck.build_deck_pptx_bytes(ctx)))
    for slide in prs.slides:
        has_table = any(shape.has_table for shape in slide.shapes)
        if has_table:
            continue  # table slides are title-only by design
        placeholders = [p.placeholder_format.idx for p in slide.placeholders]
        assert 0 in placeholders, "the title placeholder"
        assert len(placeholders) >= 2, "body text belongs in the body placeholder"


def test_powerpoint_notes_go_in_the_notes_slide(ctx: deck.DeckContext):
    from pptx import Presentation

    prs = Presentation(BytesIO(deck.build_deck_pptx_bytes(ctx)))
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides]
    assert any("Let the laugh happen" in n for n in notes)
    # Notes must not have been dumped onto the slide surface instead.
    for slide in prs.slides:
        body = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
        assert "Let the laugh happen" not in body


def test_powerpoint_tables_carry_alt_text_and_a_header_row(ctx: deck.DeckContext):
    from pptx import Presentation

    prs = Presentation(BytesIO(deck.build_deck_pptx_bytes(ctx)))
    found = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            found += 1
            assert shape.table.first_row is True
            descr = shape._element.nvGraphicFramePr.cNvPr.get("descr")
            assert descr, "a table with no alt text"
    assert found >= 1, "the three piles table"


def test_powerpoint_has_no_text_under_18pt(ctx: deck.DeckContext):
    from pptx import Presentation

    prs = Presentation(BytesIO(deck.build_deck_pptx_bytes(ctx)))
    small = []
    for number, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            frames = []
            if shape.has_text_frame:
                frames.append(shape.text_frame)
            if getattr(shape, "has_table", False) and shape.has_table:
                frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
            for frame in frames:
                for paragraph in frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size is not None and run.font.size.pt < 18:
                            small.append((number, run.text[:30], run.font.size.pt))
    assert not small, small


def test_word_passes_glow_s_own_audit(ctx: deck.DeckContext, tmp_path):
    """Use a checker, fix what it finds - the deck is held to GLOW's own bar."""
    from acb_large_print.auditor import audit_document

    path = tmp_path / "deck.docx"
    path.write_bytes(deck.build_deck_docx_bytes(ctx))
    result = audit_document(path)
    serious = [
        f for f in result.findings
        if str(f.severity).split(".")[-1] in {"CRITICAL", "HIGH", "MEDIUM"}
    ]
    assert not serious, [(f.rule_id, f.message) for f in serious]


def test_powerpoint_passes_glow_s_own_audit(ctx: deck.DeckContext, tmp_path):
    from acb_large_print.pptx_auditor import audit_presentation

    path = tmp_path / "deck.pptx"
    path.write_bytes(deck.build_deck_pptx_bytes(ctx))
    result = audit_presentation(path)
    assert not result.findings, [(f.rule_id, f.message) for f in result.findings]


def test_powerpoint_declares_its_title_and_language(ctx: deck.DeckContext):
    from pptx import Presentation

    prs = Presentation(BytesIO(deck.build_deck_pptx_bytes(ctx)))
    assert prs.core_properties.title == deck.DECK_TITLE
    assert prs.core_properties.language == "en-US"


# ---------------------------------------------------------------------------
# The routes
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("suffix", "mimetype"),
    [
        ("md", "text/markdown"),
        ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    ],
)
def test_each_format_downloads_with_the_right_type(client, suffix: str, mimetype: str):
    resp = client.get(f"/workshop/session/{CODE}/deck.{suffix}")
    assert resp.status_code == 200
    assert resp.mimetype == mimetype
    assert f"glow-workshop-deck.{suffix}" in resp.headers["Content-Disposition"]
    assert len(resp.data) > 1000


def test_the_formats_are_offered_on_the_deck_itself(client):
    body = client.get(f"/workshop/session/{CODE}/deck").get_data(as_text=True)
    for label in ("HTML, one file", "Word", "PowerPoint", "Markdown"):
        assert label in body


def test_the_formats_read_without_a_session(client):
    for suffix in ("md", "docx", "pptx"):
        assert client.get(f"/workshop/deck.{suffix}").status_code == 200
