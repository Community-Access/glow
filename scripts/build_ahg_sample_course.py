#!/usr/bin/env python
"""Build the AHG 2026 sample course: PSY 101, with its barriers planted on purpose.

Every file this writes is deliberately inaccessible. That is the point: the
workshop's agents need something real to find, and facilitators need to know
exactly what is there so they can tell whether an agent found it. The answer
key in ``docs/ahg-2026/sample-course/answer-key.md`` lists every planted
barrier; keep the two in step when you change either.

The course is fictional, the people are fictional, and every word of the
reading is original. The tone is light on purpose -- a course about memory
and procrastination that forgets things and runs late -- but nothing in it
mocks a student or a disability.

    python scripts/build_ahg_sample_course.py

Needs python-docx, python-pptx, openpyxl, Pillow and PyMuPDF, all of which the
web app already depends on.
"""

from __future__ import annotations

import io
import random
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "docs" / "ahg-2026" / "sample-course"
FONTS = Path("C:/Windows/Fonts")

COURSE = "PSY 101: Introduction to Psychology"
TERM = "Fall 2026"
PROFESSOR = "Dr. Dana Whitfield"
SCHOOL = "Mesa Ridge State University"


# ---------------------------------------------------------------------------
# Pictures (all drawn here, so nothing is borrowed)
# ---------------------------------------------------------------------------


def _font(size: int, name: str = "arial.ttf"):
    from PIL import ImageFont

    for candidate in (FONTS / name, Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")):
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def _png(image) -> bytes:
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _jpeg(image, quality: int = 55) -> bytes:
    buffer = io.BytesIO()
    image.convert("L").save(buffer, format="JPEG", quality=quality, optimize=True)
    return buffer.getvalue()


def draw_brain_logo() -> bytes:
    """A cheerful cartoon brain wearing reading glasses. No alt text, on purpose."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (480, 320), "white")
    d = ImageDraw.Draw(img)
    d.ellipse((110, 40, 370, 250), fill=(244, 170, 190), outline=(150, 60, 90), width=6)
    for x in (160, 210, 260, 310):
        d.arc((x - 40, 70, x + 40, 160), 200, 340, fill=(150, 60, 90), width=5)
        d.arc((x - 40, 130, x + 40, 220), 20, 160, fill=(150, 60, 90), width=5)
    d.ellipse((165, 120, 225, 170), outline="black", width=6)
    d.ellipse((255, 120, 315, 170), outline="black", width=6)
    d.line((225, 140, 255, 140), fill="black", width=6)
    d.text((140, 262), "PSY 101", fill=(150, 60, 90), font=_font(40, "arialbd.ttf"))
    return _png(img)


def draw_keys_chart() -> bytes:
    """A bar chart that exists only as a picture: where my keys turned up."""
    from PIL import Image, ImageDraw

    data = [("Coat pocket", 41), ("Fridge", 7), ("In the door", 23), ("My hand", 19), ("Never found", 10)]
    img = Image.new("RGB", (900, 520), "white")
    d = ImageDraw.Draw(img)
    d.text((30, 18), "Where my keys turned up, 2025 (n = 100 searches)", fill="black", font=_font(30, "arialbd.ttf"))
    base, left, width, gap = 470, 90, 120, 45
    d.line((left - 10, 80, left - 10, base), fill="black", width=3)
    d.line((left - 10, base, 880, base), fill="black", width=3)
    colors = [(46, 125, 50), (198, 40, 40), (46, 125, 50), (46, 125, 50), (198, 40, 40)]
    for i, (label, value) in enumerate(data):
        x = left + i * (width + gap)
        d.rectangle((x, base - value * 8, x + width, base), fill=colors[i])
        d.text((x, base + 8), label, fill="black", font=_font(20))
    return _png(img)


def draw_memory_diagram() -> bytes:
    """Three boxes and arrows: the multi-store model, as a picture with no description."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (900, 300), "white")
    d = ImageDraw.Draw(img)
    boxes = [("Sensory", 30), ("Short-term", 340), ("Long-term", 650)]
    for label, x in boxes:
        d.rounded_rectangle((x, 90, x + 220, 200), radius=18, outline=(30, 60, 140), width=5, fill=(225, 235, 250))
        d.text((x + 28, 128), label, fill=(30, 60, 140), font=_font(30, "arialbd.ttf"))
    for x in (250, 560):
        d.line((x + 5, 145, x + 85, 145), fill="black", width=5)
        d.polygon([(x + 85, 133), (x + 85, 157), (x + 100, 145)], fill="black")
    d.text((365, 230), "rehearsal keeps it here", fill=(120, 120, 120), font=_font(20))
    return _png(img)


# ---------------------------------------------------------------------------
# 1. Syllabus (Word)
# ---------------------------------------------------------------------------


def build_syllabus(path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    # Barrier: no document title in the properties.
    doc.core_properties.title = ""
    doc.core_properties.author = "dwhitfield"

    def fake_heading(text: str, size: int = 16) -> None:
        # Barrier: looks like a heading, is a bold Normal paragraph.
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(size)

    def add_link(paragraph, text: str, url: str) -> None:
        part = paragraph.part
        r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        rpr.append(color)
        rpr.append(underline)
        run.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        link.append(run)
        paragraph._p.append(link)

    # Barrier: a picture with no alt text, first thing on the page.
    doc.add_picture(io.BytesIO(draw_brain_logo()))

    fake_heading(f"{COURSE} -- {TERM}", 20)
    p = doc.add_paragraph("Mondays and Wednesdays, 10:00 to 11:15, Hartley Hall 104 (the room where the projector sometimes works)")

    # Barrier: a layout table, no header row, used to put two columns side by side.
    table = doc.add_table(rows=3, cols=2)
    cells = [
        ("Instructor", f"{PROFESSOR}"),
        ("Office hours", "Tuesdays 2:00 to 4:00, or whenever the coffee machine is working"),
        ("Email", "dwhitfield@mesaridge.example.edu -- I answer within two business days, or one if you include a cat picture"),
    ]
    for row, (a, b) in zip(table.rows, cells):
        row.cells[0].text = a
        row.cells[1].text = b

    doc.add_paragraph("")  # Barrier: empty paragraphs used for spacing.
    doc.add_paragraph("")

    fake_heading("Course Description")
    doc.add_paragraph(
        "Why do we remember song lyrics from 2009 but not why we walked into the kitchen? "
        "This course is a tour of how minds work: perception, memory, learning, motivation, "
        "emotion, and the surprisingly busy social life of the human brain. You will leave with "
        "a working vocabulary, a healthy suspicion of anything that says it will \"rewire your brain "
        "in seven days,\" and at least one fact you will tell people at parties whether they ask or not."
    )

    fake_heading("Required Materials")
    p = doc.add_paragraph("Readings in ")
    r = p.add_run("red")
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p.add_run(" are required. Readings in ")
    r = p.add_run("green")
    r.font.color.rgb = RGBColor(0x00, 0x80, 0x00)
    p.add_run(" are optional.")  # Barrier: meaning carried by color alone.
    for title, colour in [
        ("Chapter 1: What Psychology Is, and What It Is Not", "C00000"),
        ("Chapter 3: Memory, or, Where Did I Put That", "C00000"),
        ("A Short History of Forgetting (scanned, posted in Week 3)", "C00000"),
        ("Optional podcast: The Science of Naps", "008000"),
    ]:
        para = doc.add_paragraph()
        run = para.add_run(title)
        run.font.color.rgb = RGBColor.from_string(colour)

    fake_heading("Grading")
    # Barrier: a typed, fake numbered list.
    for line in [
        "1) Weekly reflections -- 20%",
        "2) Two quizzes -- 30%",
        "3) Lab reports -- 20%",
        "4) Final project -- 30%",
    ]:
        doc.add_paragraph(line)

    p = doc.add_paragraph()
    r = p.add_run("NO MAKEUP EXAMS WILL BE GIVEN UNDER ANY CIRCUMSTANCES WHATSOEVER.")  # Barrier: all caps, and it contradicts the accommodations statement below.
    r.bold = True

    fake_heading("Late Work")
    p = doc.add_paragraph()
    r = p.add_run(
        "Late work is accepted with a small penalty. We study procrastination in Week 9, and "
        "frankly I need the data. If something serious happens, talk to me early; I would much "
        "rather adjust a deadline than lose you for the semester."
    )
    r.italic = True  # Barrier: a long passage set in italics.

    fake_heading("Links")
    p = doc.add_paragraph("For the course calendar, ")
    add_link(p, "click here", "https://example.edu/psy101/calendar")
    p.add_run(". For the writing center, ")
    add_link(p, "click here", "https://example.edu/writing-center")
    p.add_run(".")  # Barrier: two links, both called "click here".

    fake_heading("Schedule at a Glance")
    sched = doc.add_table(rows=1, cols=3)
    sched.style = "Table Grid"
    for i, h in enumerate(("Week", "Topic", "Due")):
        sched.rows[0].cells[i].text = h  # Barrier: header row not marked as a header.
    for week, topic, due in [
        ("1", "What psychology is (and why your horoscope is not it)", "--"),
        ("3", "Memory: why you walked into the kitchen", "Reflection 1"),
        ("5", "Learning: dogs, bells, and your phone's notification sound", "Lab 1"),
        ("9", "Motivation and procrastination (on time, we hope)", "Quiz 1"),
        ("14", "Social psychology: why nobody answers in a group chat", "Final project"),
    ]:
        cells = sched.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = week, topic, due

    for _ in range(3):
        doc.add_paragraph("")

    # Barrier: the accommodations statement, last, in small light-gray type.
    p = doc.add_paragraph()
    r = p.add_run(
        "Students with disabilities who need accommodations should contact the Disability "
        "Resource Center as early as possible. Approved accommodations, including extended "
        "time and alternative exam arrangements, will be provided."
    )
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0xB0, 0xB0, 0xB0)
    r.font.highlight_color = WD_COLOR_INDEX.AUTO

    doc.save(path)


# ---------------------------------------------------------------------------
# 2. Week 3 lecture (PowerPoint)
# ---------------------------------------------------------------------------


def build_lecture(path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = ""  # Barrier: no presentation title.

    title_layout, content_layout, blank_layout = prs.slide_layouts[0], prs.slide_layouts[1], prs.slide_layouts[6]

    # Slide 1: a fine title slide, so not everything is broken.
    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = "Week 3: Memory"
    s.placeholders[1].text = "Or: why you walked into the kitchen and stood there"

    # Slide 2: no title at all; text in free-floating boxes added bottom-up.
    s = prs.slides.add_slide(blank_layout)
    for top, text in [
        (5.2, "3. Long-term memory: the attic. Huge, dusty, hard to search."),
        (3.8, "2. Short-term memory: about seven things, for about twenty seconds."),
        (2.4, "1. Sensory memory: everything, for about half a second."),
    ]:
        box = s.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.5), Inches(1.0))
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    # Barrier: no title, and reading order is 3, 2, 1.

    # Slide 3: the model as a picture with no alt text.
    s = prs.slides.add_slide(content_layout)
    s.shapes.title.text = "Memory"
    s.placeholders[1].text = " "
    s.shapes.add_picture(io.BytesIO(draw_memory_diagram()), Inches(1.5), Inches(2.2), width=Inches(10))

    # Slide 4: duplicate title, a wall of tiny text.
    s = prs.slides.add_slide(content_layout)
    s.shapes.title.text = "Memory"
    body = s.placeholders[1].text_frame
    body.text = (
        "Encoding, storage and retrieval are three separate jobs, and each can fail on its own. "
        "You can store something perfectly and still fail to retrieve it, which is why the answer "
        "arrives in the shower an hour after the quiz. Context helps retrieval: studying where you "
        "will be tested, or at least in a similar state, improves recall. Sleep consolidates. "
        "Cramming produces confidence that is not, regrettably, the same thing as memory. "
        "Spacing your study across days beats one long session, even when it feels slower, "
        "because forgetting a little between sessions makes the next retrieval do more work."
    )
    for paragraph in body.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(11)  # Barrier: small text.

    # Slide 5: the keys chart, as a picture only.
    s = prs.slides.add_slide(content_layout)
    s.shapes.title.text = "A completely scientific study"
    s.placeholders[1].text = " "
    s.shapes.add_picture(io.BytesIO(draw_keys_chart()), Inches(1.8), Inches(1.8), width=Inches(9.5))

    # Slide 6: true or false, answered by color alone.
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Myth or fact?"
    rows = [
        ("We only use ten percent of our brains", False),
        ("Sleep helps you remember what you studied", True),
        ("Goldfish have a three-second memory", False),
        ("Testing yourself beats rereading your notes", True),
    ]
    shape = s.shapes.add_table(len(rows), 1, Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.6))
    for i, (claim, is_true) in enumerate(rows):
        cell = shape.table.cell(i, 0)
        cell.text = claim
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x2E, 0x7D, 0x32) if is_true else RGBColor(0xC6, 0x28, 0x28)
        for run in cell.text_frame.paragraphs[0].runs:
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # Barrier: green means fact and red means myth, and nothing else says so.

    # Slide 7: an empty title placeholder.
    s = prs.slides.add_slide(content_layout)
    s.shapes.title.text = ""
    s.placeholders[1].text = "Questions? (I will forget them by Wednesday. Email me.)"

    prs.save(path)


# ---------------------------------------------------------------------------
# 3. Gradebook template (Excel)
# ---------------------------------------------------------------------------


def build_gradebook(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active  # Barrier: default sheet name "Sheet".
    ws["A1"] = "PSY 101 Fall 2026 -- Gradebook"
    ws.merge_cells("A1:F1")  # Barrier: merged title row.
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Reflections"
    ws.merge_cells("A2:C2")  # Barrier: merged header cells over the columns.
    ws["D2"] = "Quizzes"
    ws.merge_cells("D2:E2")
    ws["A3"], ws["B3"], ws["C3"], ws["D3"], ws["E3"], ws["F3"] = "R1", "R2", "R3", "Q1", "Q2", "Status"
    green = PatternFill("solid", fgColor="2E7D32")
    red = PatternFill("solid", fgColor="C62828")
    amber = PatternFill("solid", fgColor="F9A825")
    rows = [
        (9, 10, 8, 88, 91, green),
        (7, None, 6, 72, 65, amber),
        (10, 10, 10, 97, 99, green),
        (None, None, 5, 58, None, red),
    ]
    for r, (a, b, c, d, e, fill) in enumerate(rows, start=5):  # Barrier: blank row 4 for spacing.
        for col, value in zip("ABCDE", (a, b, c, d, e)):
            ws[f"{col}{r}"] = value
        ws[f"F{r}"].fill = fill  # Barrier: status shown by color only, the cell is empty.
    ws2 = wb.create_sheet("Sheet2")  # Barrier: a second sheet with a meaningless name.
    ws2["A1"] = "Late penalty: 5% per day, capped at 20% (see syllabus, if you can find it)"
    ws["H5"] = "Green good, red bad, yellow meh"
    ws["H5"].font = Font(color="C0C0C0", size=8)
    wb.save(path)


# ---------------------------------------------------------------------------
# 4. Required reading (scanned PDF)
# ---------------------------------------------------------------------------

READING = [
    "A Short History of Forgetting",
    "",
    "In 1885 a German psychologist named Hermann Ebbinghaus did something no",
    "sensible person would do. He memorized lists of nonsense syllables -- DAX,",
    "BOK, YAT -- and then measured, over days and weeks, how much of each list",
    "he had lost. He was both the scientist and the only subject, which is",
    "either admirable dedication or a very quiet decade.",
    "",
    "What he found has held up remarkably well. Forgetting is fastest at the",
    "start: much of a new list is gone within the first day, and the loss then",
    "slows to a gentle slide. Plotted on a graph, the drop looks like a ski slope",
    "that levels off, and it has been called the forgetting curve ever since.",
    "",
    "The more useful finding is what flattens the curve. Each time the material",
    "was reviewed, it was forgotten more slowly afterwards. Reviews spread across",
    "days did far more than the same amount of review crammed into one sitting.",
    "Modern studies call this the spacing effect, and it is one of the most",
    "reliable results in all of psychology.",
    "",
    "For students the lesson is plain, if inconvenient: an hour a day for four",
    "days beats four hours the night before. For instructors it is plainer",
    "still. A reading that a student cannot open, search, enlarge or listen to",
    "is a reading they cannot review -- and a reading nobody can review is a",
    "reading everybody forgets.",
    "",
    "Discussion questions",
    "1. Why might nonsense syllables have been a clever choice, and a limited one?",
    "2. Describe a time spacing your practice helped you learn something.",
    "3. What would make this page easier for you to review next week?",
]


def build_scanned_reading(path: Path) -> None:
    import fitz
    from PIL import Image, ImageDraw, ImageFilter

    rng = random.Random(1885)
    pages = [READING[:16], READING[16:]]
    doc = fitz.open()
    for lines in pages:
        # 150 dpi, like an office copier's "scan to email" default.
        img = Image.new("L", (1275, 1650), 246)
        d = ImageDraw.Draw(img)
        font = _font(28, "cour.ttf")
        y = 130
        for line in lines:
            d.text((110 + rng.randint(-2, 2), y), line, fill=28, font=font)
            y += 46
        # A photocopier's touch: speckle, a slight tilt, a soft blur.
        for _ in range(1500):
            x, yy = rng.randint(0, 1274), rng.randint(0, 1649)
            img.putpixel((x, yy), rng.randint(120, 200))
        img = img.rotate(0.8, fillcolor=246).filter(ImageFilter.GaussianBlur(0.6))
        page = doc.new_page(width=612, height=792)
        page.insert_image(page.rect, stream=_jpeg(img))
    doc.set_metadata({})  # Barrier: no title, no language, no text layer at all.
    doc.save(path)


# ---------------------------------------------------------------------------
# 5. Lab handout (untagged PDF with an unlabeled form)
# ---------------------------------------------------------------------------


def build_lab_handout(path: Path) -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 80), "Lab 1: The Stroop Test (or, Why Your Brain Argues With Itself)", fontsize=15, fontname="hebo")
    body = (
        "Name the COLOR of each word below, not the word. Time yourself. Then do the second list.\n"
        "Most people are slower when the word and the color disagree. That delay is your brain\n"
        "refereeing a fight between reading, which is automatic, and naming, which is not."
    )
    page.insert_text((72, 115), body, fontsize=10.5, fontname="helv")
    words = [("RED", (0, 0.5, 0)), ("BLUE", (0.8, 0, 0)), ("GREEN", (0, 0, 0.8)), ("YELLOW", (0.5, 0, 0.5))]
    x = 72
    for word, colour in words:
        page.insert_text((x, 200), word, fontsize=20, fontname="hebo", color=colour)
        x += 120
    # Barrier: the exercise depends entirely on seeing color, with no alternative.

    # A drawn table with no structure behind it.
    top, left, col_w, row_h = 250, 72, 156, 24
    headers = ["Trial", "Matching (seconds)", "Mismatched (seconds)"]
    for r in range(5):
        for c in range(3):
            rect = fitz.Rect(left + c * col_w, top + r * row_h, left + (c + 1) * col_w, top + (r + 1) * row_h)
            page.draw_rect(rect, color=(0, 0, 0), width=0.7)
            text = headers[c] if r == 0 else (str(r) if c == 0 else "")
            if text:
                page.insert_text((rect.x0 + 6, rect.y0 + 16), text, fontsize=10, fontname="helv")

    # Barrier: form fields with no labels or tooltips.
    for i, y in enumerate((420, 460, 500)):
        widget = fitz.Widget()
        widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
        widget.field_name = f"Text{i + 1}"
        widget.rect = fitz.Rect(250, y, 540, y + 22)
        page.add_widget(widget)
    page.insert_text((72, 436), "_____________", fontsize=10, fontname="helv")
    page.insert_text((72, 476), "_____________", fontsize=10, fontname="helv")
    page.insert_text((72, 516), "_____________", fontsize=10, fontname="helv")

    page.insert_text((72, 720), "Hand in by Friday. Late labs accepted (see Week 9, procrastination).", fontsize=7, fontname="helv", color=(0.65, 0.65, 0.65))
    doc.set_metadata({})
    doc.save(path)


# ---------------------------------------------------------------------------
# 6. Course announcement (web page)
# ---------------------------------------------------------------------------

ANNOUNCEMENT = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>Announcement</title>
<style>
  body { font-family: Arial, sans-serif; margin: 2rem; color: #333; }
  .soft { color: #b5b5b5; }
  .due { color: #d32f2f; }
  .btn { display: inline-block; background: #7cb342; color: #ffffff; padding: 0.4rem 0.9rem; cursor: pointer; }
  .tiny { font-size: 11px; }
</style>
</head>
<body>
<div class="header"><b style="font-size:28px">PSY 101 Announcements</b></div>

<h1>Week 3 is here, and so is memory</h1>
<img src="brain.png" width="160">
<p>Hello everyone! This week we finally answer the question that has haunted humanity since the invention of kitchens:
why do you walk into a room and forget why you came in? (Short answer: doorways. Long answer: Wednesday's lecture.)</p>

<h4>What to do this week</h4>
<p class="soft">Read the scanned chapter, then do the Stroop lab. Items in red are due Friday.</p>
<ul>
  <li class="due">Reading: A Short History of Forgetting</li>
  <li>Optional: the nap podcast (research, technically)</li>
  <li class="due">Lab 1: the Stroop test</li>
</ul>
<p>Slides are posted. <a href="psy101-week3-lecture.pptx">Click here</a>. The syllabus has changed slightly. <a href="psy101-syllabus.docx">Read more</a>.</p>

<div class="btn" onclick="alert('Submitted!')">Submit reflection</div>

<h4>Quiz reminder</h4>
<p>Quiz 1 opens in Week 9. Yes, the procrastination week. Yes, on purpose.</p>

<p class="tiny soft">Need an accommodation? Contact the Disability Resource Center.</p>
</body>
</html>
"""


# ---------------------------------------------------------------------------
# 7. Lecture captions (auto-generated, uncorrected)
# ---------------------------------------------------------------------------

CAPTIONS = """WEBVTT

00:00:01.000 --> 00:00:04.500
okay so today we're talking about memory

00:00:04.500 --> 00:00:09.000
and specifically about a guy named herman ebb in house

00:00:09.000 --> 00:00:13.500
who memorized nonsense syllables like docks and bach

00:00:13.500 --> 00:00:18.000
so as you can see on this graph it drops really fast and then levels off

00:00:18.000 --> 00:00:21.500
yes in the back

00:00:21.500 --> 00:00:25.000
is this on the quiz

00:00:25.000 --> 00:00:30.000
everything is on the quiz [LAUGHTER] no it is the spacing effect that's on the quiz

00:00:30.000 --> 00:00:35.000
space your studying across days it beats cramming every time
"""


# ---------------------------------------------------------------------------


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    builders = [
        ("psy101-syllabus.docx", build_syllabus),
        ("psy101-week3-lecture.pptx", build_lecture),
        ("psy101-gradebook.xlsx", build_gradebook),
        ("psy101-reading-forgetting-scanned.pdf", build_scanned_reading),
        ("psy101-lab1-stroop.pdf", build_lab_handout),
    ]
    for name, build in builders:
        build(OUT / name)
        print(f"  {name}")
    (OUT / "psy101-announcement.html").write_text(ANNOUNCEMENT, encoding="utf-8")
    (OUT / "brain.png").write_bytes(draw_brain_logo())
    (OUT / "psy101-week3-captions.vtt").write_text(CAPTIONS, encoding="utf-8")
    print("  psy101-announcement.html, brain.png, psy101-week3-captions.vtt")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
