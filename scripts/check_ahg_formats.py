#!/usr/bin/env python
"""Check every AHG 2026 participant document in all four formats.

What a parser can see, for each source in ``scripts/ahg_formats.py``:

* the HTML, Word and PowerPoint files exist, and were built from the
  Markdown as it is now (``docs/ahg-2026/formats.json``);
* HTML: a language, a title, one h1, no skipped heading level, a scope on
  every header cell, alt text on every image, no vague link text;
* Word: title and language on the file, one Heading 1 and no skipped level,
  every table's header row marked to repeat and the table given alt text;
* PowerPoint: title and language on the file, a title on every slide, every
  title unique, the title first in reading order, every table and picture
  given alt text and every table a marked header row;
* none of the four carries an emoji.

It does not replace Word's and PowerPoint's Accessibility Checkers or a
screen reader; ``golden-readiness.md`` lists those as checks for a person.

    python scripts/check_ahg_formats.py      # exit 1 on any failure
"""

from __future__ import annotations

import json
import re
import sys
import zipfile
from html.parser import HTMLParser
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import ahg_formats as F  # noqa: E402

EMOJI = re.compile("[\U0001F300-\U0001FAFF\U00002600-\U000027BF\U0001F000-\U0001F2FF\U0000FE0F]")
VAGUE = {"click here", "here", "link", "this link", "more", "read more", "this"}


class _Html(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.lang = ""
        self.title = ""
        self.headings: list[int] = []
        self.problems: list[str] = []
        self._in_title = False
        self._link: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "html":
            self.lang = a.get("lang") or ""
        elif tag == "title":
            self._in_title = True
        elif re.fullmatch(r"h[1-6]", tag):
            self.headings.append(int(tag[1]))
        elif tag == "th" and a.get("scope") not in ("col", "row"):
            self.problems.append("a header cell has no scope")
        elif tag == "img" and a.get("alt") is None:
            self.problems.append("an image has no alt attribute")
        elif tag == "a":
            self._link = []

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        elif tag == "a" and self._link is not None:
            text = " ".join("".join(self._link).split()).lower()
            if not text or text in VAGUE:
                self.problems.append(f"vague link text: {text!r}")
            self._link = None

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        if self._link is not None:
            self._link.append(data)


def _levels(levels: list[int], what: str) -> list[str]:
    out = []
    if levels.count(1) != 1:
        out.append(f"{what}: {levels.count(1)} top-level headings, not one")
    if levels and levels[0] != 1:
        out.append(f"{what}: the first heading is level {levels[0]}")
    for before, after in zip(levels, levels[1:]):
        if after > before + 1:
            out.append(f"{what}: heading level {before} is followed by {after}")
            break
    return out


def check_html(path: Path) -> list[str]:
    text = path.read_text(encoding="utf-8")
    parser = _Html()
    parser.feed(text)
    out = list(parser.problems)
    if not parser.lang:
        out.append("no lang on <html>")
    if not parser.title.strip():
        out.append("no <title>")
    out += _levels(parser.headings, "HTML")
    if EMOJI.search(text):
        out.append("an emoji")
    return out


def check_docx(path: Path) -> list[str]:
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    out = []
    if not doc.core_properties.title:
        out.append("no title in the file properties")
    if not doc.core_properties.language:
        out.append("no language in the file properties")
    levels = []
    for p in doc.paragraphs:
        m = re.fullmatch(r"Heading (\d)", p.style.name or "")
        if m:
            levels.append(int(m.group(1)))
        if EMOJI.search(p.text):
            out.append("an emoji")
    out += _levels(levels, "Word")
    for n, table in enumerate(doc.tables, start=1):
        tr_pr = table.rows[0]._tr.trPr
        if tr_pr is None or tr_pr.find(qn("w:tblHeader")) is None:
            out.append(f"table {n}: header row not marked")
        if table._tbl.tblPr.find(qn("w:tblCaption")) is None:
            out.append(f"table {n}: no alt text")
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            out.append("contains an image; check its alt text")
    return out


def check_pptx(path: Path) -> list[str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    if not prs.core_properties.title:
        out.append("no title in the file properties")
    if not prs.core_properties.language:
        out.append("no language in the file properties")
    titles: list[str] = []
    for n, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title
        if title is None or not title.text_frame.text.strip():
            out.append(f"slide {n}: no title")
            continue
        titles.append(title.text_frame.text.strip())
        shapes = list(slide.shapes)
        if shapes and shapes[0].shape_id != title.shape_id:
            out.append(f"slide {n}: the title is not first in reading order")
        for shape in shapes:
            descr = shape._element.xpath("./*[1]/p:cNvPr/@descr")
            if shape.has_table:
                if not descr:
                    out.append(f"slide {n}: a table has no alt text")
                if not shape.table.first_row:
                    out.append(f"slide {n}: a table has no header row")
            if shape.shape_type == 13 and not descr:  # picture
                out.append(f"slide {n}: a picture has no alt text")
            if shape.has_text_frame and EMOJI.search(shape.text_frame.text):
                out.append(f"slide {n}: an emoji")
    dupes = sorted({t for t in titles if titles.count(t) > 1})
    if dupes:
        out.append(f"slide titles used twice: {dupes}")
    return out


def check_all() -> dict[str, list[str]]:
    results: dict[str, list[str]] = {}
    manifest = json.loads(F.MANIFEST.read_text(encoding="utf-8"))["documents"] if F.MANIFEST.is_file() else {}
    for rel in F.SOURCES:
        source = F.AHG / rel
        markdown = source.read_text(encoding="utf-8")
        problems = []
        if EMOJI.search(markdown):
            problems.append(f"{rel}: an emoji")
        recorded = manifest.get(rel, {}).get("sha256")
        if recorded != F.source_hash(markdown):
            problems.append(f"{rel}: changed since its formats were built; run python scripts/build_ahg_kit.py --formats-only")
        for out in F.outputs_for(rel):
            path = F.AHG / out
            if not path.is_file():
                problems.append(f"{out}: missing")
                continue
            check = {".html": check_html, ".docx": check_docx, ".pptx": check_pptx}[path.suffix]
            problems += [f"{out}: {p}" for p in check(path)]
        results[rel] = problems
    return results


def main() -> int:
    results = check_all()
    failed = {k: v for k, v in results.items() if v}
    for problems in failed.values():
        for p in problems:
            print(p)
    print(f"{len(results)} documents, {len(results) * 4} files: "
          f"{len(results) - len(failed)} pass, {len(failed)} fail")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
