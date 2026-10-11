#!/usr/bin/env python
"""Check that every sample course file opens and every planted barrier is still there.

Test cases Mat-01 and Mat-02 in ahg.md. The answer key in
``docs/ahg-2026/answer-key.md`` promises facilitators these barriers exist;
if ``scripts/build_ahg_sample_course.py`` changes, run this again.

    python scripts/check_ahg_sample_course.py

Exit code 0 when every check passes.
"""
import re
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

C = Path(__file__).resolve().parents[1] / "docs" / "ahg-2026" / "kit" / "sample-course"
checks = []


def check(name, ok):
    checks.append((name, ok))
    print(f"{'Pass' if ok else 'FAIL'}  {name}")


# Syllabus
d = Document(C / "psy101-syllabus.docx")
paras = d.paragraphs
check("syllabus opens", True)
check("syllabus: no heading styles used", not any(p.style.name.startswith("Heading") for p in paras))
check("syllabus: layout table without header", len(d.tables) >= 2 and d.tables[0].rows[0].cells[0].text == "Instructor")
reds = [r for p in paras for r in p.runs if r.font.color and r.font.color.rgb is not None and str(r.font.color.rgb) == "C00000"]
check("syllabus: required readings shown only in red", len(reds) >= 3)
check("syllabus: typed '1)' list", sum(p.text.startswith(("1)", "2)", "3)", "4)")) for p in paras) == 4)
check("syllabus: NO MAKEUP EXAMS in capitals", any("NO MAKEUP EXAMS" in p.text for p in paras))
check("syllabus: italic passage", any(r.italic for p in paras for r in p.runs))
check("syllabus: two 'click here' links", d.element.xml.count(">click here<") == 2)
check("syllabus: empty spacer paragraphs", sum(1 for p in paras if not p.text.strip()) >= 4)
last = [p for p in paras if p.text.strip()][-1]
check("syllabus: accommodations statement last, small and gray",
      "Disability Resource Center" in last.text and last.runs[0].font.size.pt <= 8 and str(last.runs[0].font.color.rgb) == "B0B0B0")
check("syllabus: no document title", not d.core_properties.title)
check("syllabus: schedule header row not marked", "w:tblHeader" not in d.tables[1]._tbl.xml)

# Lecture
prs = Presentation(C / "psy101-week3-lecture.pptx")
slides = list(prs.slides)
check("lecture opens", True)
check("lecture: slide 2 has no title", slides[1].shapes.title is None)
tops = [round(s.top) for s in slides[1].shapes]
check("lecture: slide 2 reading order bottom-up", tops == sorted(tops, reverse=True))
pics = [sh for s in slides for sh in s.shapes if sh.shape_type == 13]
check("lecture: pictures carry 'image.png' as alt text", pics and all(p._element.nvPicPr.cNvPr.get("descr") == "image.png" for p in pics))
titles = [s.shapes.title.text for s in slides if s.shapes.title is not None]
check("lecture: duplicate 'Memory' titles", titles.count("Memory") == 2)
check("lecture: empty title on last slide", slides[-1].shapes.title is not None and not slides[-1].shapes.title.text.strip())
tbl = [sh for sh in slides[5].shapes if sh.has_table][0].table
fills = {str(tbl.cell(i, 0).fill.fore_color.rgb) for i in range(4)}
check("lecture: myth or fact by red and green only", fills == {"2E7D32", "C62828"} and not any(w in tbl.cell(0, 0).text.lower() for w in ("myth", "fact")))
small = [r.font.size.pt for sh in slides[3].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
check("lecture: 11pt wall of text", small and max(small) <= 11)

# Gradebook
wb = load_workbook(C / "psy101-gradebook.xlsx")
ws = wb.worksheets[0]
check("gradebook opens", True)
check("gradebook: merged title and headers", {"A1:F1", "A2:C2"} <= {str(r) for r in ws.merged_cells.ranges})
check("gradebook: default sheet names", [s.title for s in wb.worksheets] == ["Sheet", "Sheet2"])
check("gradebook: status only by fill", all(ws[f"F{r}"].value is None and ws[f"F{r}"].fill.fgColor.rgb not in (None, "00000000") for r in range(5, 9)))
check("gradebook: blank spacer row 4", all(ws[f"{c}4"].value is None for c in "ABCDEF"))
check("gradebook: tiny gray legend", ws["H5"].font.size == 8 and "meh" in ws["H5"].value)

# PDFs
scan = fitz.open(C / "psy101-reading-forgetting-scanned.pdf")
check("reading opens", True)
check("reading: no text layer on any page", all(not p.get_text().strip() for p in scan) and all(p.get_images() for p in scan))
lab = fitz.open(C / "psy101-lab1-stroop.pdf")
widgets = [w for p in lab for w in p.widgets()]
check("lab opens", True)
check("lab: three unlabeled fields", len(widgets) == 3 and all(not (w.field_label or "").strip() for w in widgets))
check("lab: color words in mismatched colors", "RED" in lab[0].get_text() and "GREEN" in lab[0].get_text())
check("lab: untagged", not lab.pdf_catalog() or "StructTreeRoot" not in lab.xref_object(lab.pdf_catalog()))

# Web page and captions
html = (C / "psy101-announcement.html").read_text(encoding="utf-8")
check("announcement: no lang", re.search(r"<html(?![^>]*lang)", html) is not None)
check("announcement: bold div as heading", '<div class="header"><b' in html)
check("announcement: h1 then h4", "<h1>" in html and "<h4>" in html and "<h2>" not in html and "<h3>" not in html)
check("announcement: image without alt", re.search(r"<img(?![^>]*alt=)", html) is not None)
check("announcement: click here and read more", "Click here" in html and "Read more" in html)
check("announcement: clickable div button", '<div class="btn" onclick=' in html)
check("announcement: items in red", "Items in red" in html)
vtt = (C / "psy101-week3-captions.vtt").read_text(encoding="utf-8")
check("captions: misheard name", "herman ebb in house" in vtt and "docks and bach" in vtt)
check("captions: no speaker labels", "is this on the quiz" in vtt and ">>" not in vtt)
check("captions: undescribed graph", "as you can see on this graph" in vtt)

print(f"{sum(ok for _, ok in checks)} of {len(checks)} pass")
raise SystemExit(0 if all(ok for _, ok in checks) else 1)
