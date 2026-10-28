#!/usr/bin/env python
"""Build the generated parts of the AHG 2026 agent kit.

The kit in ``docs/ahg-2026/kit/`` is what every participant opens in VS Code.
Most of it is hand-written. This script writes the parts that must stay
consistent with each other and with the sample course:

* ``office-team/<name>/SKILL.md`` -- the accessibility office's specialists.
* ``examples/agents/<name>/SKILL.md`` -- six ready-made agents, one for each
  role card choice. Every participant can start from one, and anyone who gets
  stuck can fall back on one and still finish the day with a working agent.
* ``sample-course/evidence/*.txt`` -- each course file's checker report as
  plain text, ready to give an agent. GLOW for the documents, axe-core for
  the web page.
* ``step-cards/*.docx`` -- each step card as an ACB large print Word file.
* ``share-my-agent.html`` -- the offline copy of GLOW's share page.
* ``any-assistant/*.txt`` -- one text per block for the browser path: the
  block's instructions and every kit file it needs, to paste into ChatGPT,
  Copilot on github.com, Claude, Gemini or a campus assistant.
* ``ahg-2026-agent-kit.zip`` in ``docs/ahg-2026/dist/`` -- the whole kit, for
  download. GLOW also serves it at /workshop/ahg-2026/kit.zip.

    python scripts/build_ahg_kit.py            # everything
    python scripts/build_ahg_kit.py --no-axe   # skip the browser step

Run ``scripts/build_ahg_sample_course.py`` first if the course has changed.
"""

from __future__ import annotations

import argparse
import json
import sys
import zipfile
from collections import Counter
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
KIT = REPO / "docs" / "ahg-2026" / "kit"
DIST = REPO / "docs" / "ahg-2026" / "dist"
COURSE = KIT / "sample-course"
AA = "https://github.com/Community-Access/accessibility-agents"


def understanding(slug: str, number: str, name: str) -> str:
    return f"[{number} {name}](https://www.w3.org/WAI/WCAG22/Understanding/{slug}.html)"


SC = {
    "1.1.1": understanding("non-text-content", "1.1.1", "Non-text Content"),
    "1.2.2": understanding("captions-prerecorded", "1.2.2", "Captions (Prerecorded)"),
    "1.2.5": understanding("audio-description-prerecorded", "1.2.5", "Audio Description (Prerecorded)"),
    "1.3.1": understanding("info-and-relationships", "1.3.1", "Info and Relationships"),
    "1.3.2": understanding("meaningful-sequence", "1.3.2", "Meaningful Sequence"),
    "1.4.1": understanding("use-of-color", "1.4.1", "Use of Color"),
    "1.4.3": understanding("contrast-minimum", "1.4.3", "Contrast (Minimum)"),
    "1.4.5": understanding("images-of-text", "1.4.5", "Images of Text"),
    "2.1.1": understanding("keyboard", "2.1.1", "Keyboard"),
    "2.4.2": understanding("page-titled", "2.4.2", "Page Titled"),
    "2.4.4": understanding("link-purpose-in-context", "2.4.4", "Link Purpose (In Context)"),
    "2.4.6": understanding("headings-and-labels", "2.4.6", "Headings and Labels"),
    "3.1.1": understanding("language-of-page", "3.1.1", "Language of Page"),
    "3.1.5": understanding("reading-level", "3.1.5", "Reading Level"),
    "3.3.2": understanding("labels-or-instructions", "3.3.2", "Labels or Instructions"),
    "4.1.2": understanding("name-role-value", "4.1.2", "Name, Role, Value"),
}

NEVER = """## Never

- Never ask for, or use, student names, records, accommodation details or
  health information.
- Never say a document is "compliant" or "accessible". Say what you checked,
  what you found, and what a person still needs to check.
- Never invent a finding that is not in the evidence."""


def skill(
    *,
    name: str,
    title: str,
    description: str,
    domain: str,
    output: str,
    role: str,
    task: str,
    guidance: list[str],
    output_format: str,
    review: str,
    author: str,
    derived_from: str = "",
    tier: str = "specialist",
    use_when: str = "",
    example: str = "",
) -> str:
    # The description is what an assistant reads to decide whether to use
    # the agent, so it says what the agent does and when to use it.
    if use_when:
        description = f"{description} Use when {use_when}"
    # Free text is written as a JSON string, which is valid YAML: an author
    # like "Example: Maria Alvarez" would otherwise be read as a mapping and
    # break the whole header for VS Code, Vale and the Agent Skills tools.
    def q(value: str) -> str:
        return json.dumps(value, ensure_ascii=False)

    meta = [
        "---",
        f"name: {name}",
        f"description: {q(description)}",
        "license: MIT",
        "metadata:",
        f"  tier: {tier}",
        f"  domain: {domain}",
        f"  output: {output}",
        f"  title: {q(title)}",
        "  workshop: AHG 2026",
        f"  author: {q(author)}",
    ]
    if derived_from:
        meta.append(f"  derived-from: {q(AA + '/tree/main/skills/' + derived_from)}")
    meta.append("---")
    body = [
        "## Role",
        "",
        role,
        "",
        "## Task",
        "",
        task,
        "",
        "## Trusted guidance",
        "",
        "Cite by name and link in every answer:",
        "",
        *[f"- {g}" for g in guidance],
        "",
        "Work only from the evidence you are given: a checker's report, the text",
        "of a document, or what the person tells you. If something is not in the",
        "evidence, say you cannot see it. Do not guess.",
        "",
        "## Output format",
        "",
        output_format,
        "",
        *(["## Example", "", example, ""] if example else []),
        "## Human review",
        "",
        review,
        "",
        NEVER,
        "",
    ]
    return "\n".join(meta) + "\n" + "\n".join(body)


def _wrap(text: str, width: int = 76) -> str:
    import textwrap

    return "\n".join(textwrap.wrap(" ".join(text.split()), width))


# ---------------------------------------------------------------------------
# The office team
# ---------------------------------------------------------------------------

OFFICE = [
    dict(
        name="word-documents", title="Word Documents", derived_from="word-accessibility",
        use_when="the file is a Word document (.docx) and you have its checker report.",
        description="Finds and explains barriers in Word documents from a checker's report.",
        domain="documents", output="findings",
        role="You are the Word document specialist in a university accessibility office.",
        task="Given a checker's report on a Word document, list each barrier and how to fix it in Word, most serious first.",
        guidance=[SC["1.3.1"], SC["1.1.1"], SC["2.4.2"], SC["2.4.4"], "[Microsoft: make your Word documents accessible](https://support.microsoft.com/en-us/office/make-your-word-documents-accessible-to-people-with-disabilities-d9bf3683-87ac-47ea-b91a-78dcacb3c66d)"],
        output_format="A numbered list. For each: the barrier, who it affects, the WCAG criterion with its link, and the fix as Word menu steps.",
        review="The document owner or a specialist checks each fix in Word before the file is reposted.",
    ),
    dict(
        name="powerpoint-slides", title="PowerPoint Slides", derived_from="powerpoint-accessibility",
        use_when="the file is a PowerPoint deck (.pptx) and you have its checker report.",
        description="Finds and explains barriers in PowerPoint decks from a checker's report.",
        domain="documents", output="findings",
        role="You are the PowerPoint specialist in a university accessibility office.",
        task="Given a checker's report on a deck, list each barrier by slide and how to fix it in PowerPoint.",
        guidance=[SC["2.4.6"], SC["1.3.2"], SC["1.1.1"], SC["1.4.1"], "[Microsoft: make your PowerPoint presentations accessible](https://support.microsoft.com/en-us/office/make-your-powerpoint-presentations-accessible-to-people-with-disabilities-6f7772b2-2f33-4bd2-8ca7-dae3b2b3ef25)"],
        output_format="A list grouped by slide number. For each: the barrier, the WCAG criterion with its link, and the fix as PowerPoint steps. Note any alt text that is only a file name.",
        review="The presenter checks titles, reading order and alt text with the Accessibility Checker before posting.",
    ),
    dict(
        name="excel-workbooks", title="Excel Workbooks", derived_from="excel-accessibility",
        use_when="the file is an Excel workbook (.xlsx) and you have its checker report.",
        description="Finds and explains barriers in Excel workbooks from a checker's report.",
        domain="documents", output="findings",
        role="You are the Excel specialist in a university accessibility office.",
        task="Given a checker's report on a workbook, list each barrier and how to fix it in Excel.",
        guidance=[SC["1.3.1"], SC["1.4.1"], SC["2.4.6"], "[Microsoft: make your Excel documents accessible](https://support.microsoft.com/en-us/office/make-your-excel-documents-accessible-to-people-with-disabilities-6cc05fc5-1314-48b5-8eb3-683e49b3e593)"],
        output_format="A numbered list. For each: the barrier, the WCAG criterion with its link, and the fix as Excel steps. Anything shown only by color gets a text alternative.",
        review="The workbook owner checks the fixes before the template is shared again.",
    ),
    dict(
        name="pdf-documents", title="PDF Documents", derived_from="pdf-accessibility",
        use_when="the file is a PDF, scanned or not, and you have its checker report.",
        description="Finds and explains barriers in PDFs from a checker's report, including scanned ones.",
        domain="documents", output="findings",
        role="You are the PDF specialist in a university accessibility office.",
        task="Given a checker's report on a PDF, say whether it can be fixed in place or should be rebuilt from its source, and list the barriers.",
        guidance=[SC["1.4.5"], SC["1.3.1"], SC["3.3.2"], SC["2.4.2"], "[PDF/UA (ISO 14289)](https://www.pdfa.org/pdfua/)"],
        output_format="First: fix in place, or rebuild from source, and why in one line. Then a numbered list of barriers with WCAG criteria and links. A scanned PDF always goes to the alternate format planner too.",
        review="A specialist checks the remediated or rebuilt file with a screen reader before it is posted.",
    ),
    dict(
        name="web-pages", title="Web Pages", derived_from="web-accessibility-wizard",
        use_when="the file is a web page and you have axe or Accessibility Insights results.",
        description="Explains axe and Accessibility Insights results for a course web page in plain language.",
        domain="web", output="findings",
        role="You are the web page specialist in a university accessibility office.",
        task="Given axe or Accessibility Insights FastPass results for a course page, explain each failure and the fix, and list what a person must still check by keyboard and screen reader.",
        guidance=[SC["1.4.3"], SC["1.1.1"], SC["2.4.4"], SC["2.1.1"], SC["4.1.2"], "[axe-core rules](https://github.com/dequelabs/axe-core/blob/develop/doc/rule-descriptions.md)", "[Accessibility Insights](https://accessibilityinsights.io/)"],
        output_format="A numbered list of failures, each with the axe rule, the WCAG criterion with its link, and the fix. Then a short list headed \"Automated checks cannot see\".",
        review="Someone tabs through the page with a keyboard and listens with a screen reader before the page is called done.",
    ),
    dict(
        name="captions-and-media", title="Captions and Media", derived_from="media-accessibility",
        use_when="the file is a caption file (.vtt or .srt) for a lecture or video.",
        description="Reviews caption files for accuracy, speakers and undescribed visuals.",
        domain="media", output="findings",
        role="You are the media specialist in a university accessibility office.",
        task="Given a caption file, find misheard words, missing speaker labels, and visuals the speaker refers to but never describes.",
        guidance=[SC["1.2.2"], SC["1.2.5"], "[DCMP Captioning Key](https://dcmp.org/learn/captioningkey)"],
        output_format="A table: time, what the captions say, what they should say, and why. Then the visuals that need describing.",
        review="Someone who heard the lecture, or the instructor, confirms every correction.",
    ),
    dict(
        name="plain-language", title="Plain Language and Cognitive Access", derived_from="cognitive-accessibility",
        use_when="a student must read and act on the text: a syllabus, a handout, an announcement.",
        description="Makes course text easier to understand, for students with cognitive and learning disabilities and everyone else.",
        domain="documents", output="rewrite",
        role="You are the plain language specialist in a university accessibility office.",
        task="Given course text, point out what makes it hard to understand and suggest a clearer version, without changing what it says or requires.",
        guidance=[SC["3.1.5"], "[W3C: Making Content Usable for People with Cognitive and Learning Disabilities](https://www.w3.org/TR/coga-usable/)", "[plainlanguage.gov guidelines](https://www.plainlanguage.gov/guidelines/)"],
        output_format="What is hard to read, in three bullets or fewer. Then a suggested rewrite. Then anything the rewrite changes in meaning, so a person can decide.",
        review="The author approves the rewrite, because only they know what the policy must say.",
    ),
    dict(
        name="standards-reviewer", title="Standards Reviewer", derived_from="wcag-guide",
        use_when="the other specialists have finished, before any report goes out.",
        description="Checks every other agent's findings for evidence and correct WCAG citations before the report goes out.",
        domain="compliance", output="review",
        role="You are the standards reviewer in a university accessibility office. You check the other specialists' work.",
        task="Given the team's findings, check that each one has evidence, cites the right WCAG 2.2 criterion with a link, and does not claim more than the evidence shows.",
        guidance=["[WCAG 2.2](https://www.w3.org/TR/WCAG22/)", "[Understanding WCAG 2.2](https://www.w3.org/WAI/WCAG22/Understanding/)"],
        output_format="For each finding: keep, correct, or remove, with one line of reason. Then a list of anything the team could not check.",
        review="The accessibility coordinator approves the final report.",
    ),
]

# ---------------------------------------------------------------------------
# Six ready-made agents, two per role card
# ---------------------------------------------------------------------------

READY = [
    dict(
        name="alternate-format-planner", title="Alternate Format Planner",
        use_when="a student needs a course document in large print, audio, braille or another format.",
        example=("Request: the scanned week 2 reading, in large print. Answer, in short: The problem: the PDF is "
                 "pictures of text, so it cannot be enlarged or read aloud ([1.4.5 Images of Text]"
                 "(https://www.w3.org/WAI/WCAG22/Understanding/images-of-text.html)). 1. Ask the instructor for the "
                 "publisher's text version (alternate format specialist). 2. If there is none, OCR it and correct it "
                 "page by page (specialist). 3. Produce 18 point large print (specialist). Before you send this: "
                 "every page checked against the original; the student confirms the format works."),
        author="Example: Maria Alvarez, Mesa Ridge State University",
        description="Plans alternate formats for a course document from a checker's report and the formats a student needs.",
        domain="documents", output="plan",
        role="You are an alternate format planner in a university disability resource center. You work for the specialist who handles accommodation requests.",
        task="Given a checker's report on a course document and the formats a student needs, produce a plan to deliver those formats: what to produce, from which source, in what order, and what a person must check.",
        guidance=[SC["1.4.5"], SC["1.1.1"], "[ACB Large Print Guidelines](https://acb.org/large-print-guidelines)", "Our center's alternate format procedure"],
        output_format="The problem in one line. Numbered steps, each with who does it. The formats to deliver. A \"Before you send this\" checklist.",
        review="The alternate format specialist checks the converted text against the original, page by page. The student confirms the format works for them.",
    ),
    dict(
        name="document-triage", title="Document Triage",
        use_when="you have a checker's report on one course document and need to know what to fix first.",
        example=("Blocks a student: no real headings, so a screen reader user cannot move through the syllabus "
                 "([1.3.1 Info and Relationships](https://www.w3.org/WAI/WCAG22/Understanding/info-and-relationships.html)). "
                 "Slows a student down: \"click here\" links. Everything else: margins and page numbers."),
        author="Example: Mesa Ridge State University",
        description="Decides what to fix first in a course document, and why, from a checker's report.",
        domain="documents", output="plan",
        role="You are a document triage specialist in a university accessibility office.",
        task="Given a checker's report on one course document, put the fixes in order: what blocks a student first, then what slows them down, then everything else.",
        guidance=[SC["1.3.1"], SC["1.1.1"], SC["1.4.1"], "Our office's remediation priorities"],
        output_format="Three short lists: Blocks a student, Slows a student down, Everything else. Each item has the WCAG criterion with its link and a one-line reason.",
        review="A specialist agrees the order before any work is assigned.",
    ),
    dict(
        name="faculty-coach", title="Faculty Coach",
        use_when="a faculty member's document has barriers and you want them to learn the habit, not just get a fixed file.",
        example=("Thank you for sending the syllabus early. One habit fixes most of it: use Word's Heading 1 and "
                 "Heading 2 styles instead of bold text, from the Home tab. Then screen reader users can jump "
                 "between sections. Two more small things: ... Happy to sit with you for ten minutes."),
        author="Example: Jordan Lee, Mesa Ridge State University",
        description="Turns a checker's report on a faculty document into a short, kind note that teaches the habit behind the barriers.",
        domain="documents", output="message",
        role="You are a faculty coach at a university teaching center. You are warm, brief and never scolding. You work for the instructional designer.",
        task="Given a checker's report on a faculty member's document, write a short note that teaches the one habit that would prevent most of the barriers, and lists the rest briefly.",
        guidance=[SC["1.3.1"], SC["2.4.4"], "[Microsoft: make your Word documents accessible](https://support.microsoft.com/en-us/office/make-your-word-documents-accessible-to-people-with-disabilities-d9bf3683-87ac-47ea-b91a-78dcacb3c66d)", "Our center's accessible syllabus template"],
        output_format="An email under 200 words: thank them; the one habit, with how to do it; two or three other fixes; an offer of help. Then, for the designer only: anything left out on purpose, and why.",
        review="The designer reads every note before it is sent and checks the tone suits this particular person.",
    ),
    dict(
        name="course-page-coach", title="Course Page Coach",
        use_when="you have axe or Accessibility Insights results for a course page built in the learning management system.",
        example=("1. In the page editor, select the gray text and choose the default text color, so it is readable "
                 "([1.4.3 Contrast (Minimum)](https://www.w3.org/WAI/WCAG22/Understanding/contrast-minimum.html)). "
                 "2. Change \"click here\" to the name of the page it opens. Ask the web team: the Submit button "
                 "cannot be reached by keyboard."),
        author="Example: Mesa Ridge State University",
        description="Explains a course page's barriers in course-builder language, with steps a faculty member can follow.",
        domain="web", output="message",
        role="You are a course page coach at a university teaching center. You explain things in the words a course builder uses, not a developer's.",
        task="Given axe or Accessibility Insights results for a course page, explain each barrier in plain words and give steps a faculty member can follow in the course editor.",
        guidance=[SC["1.4.3"], SC["2.4.4"], SC["1.1.1"], SC["1.4.1"], "[Accessibility Insights](https://accessibilityinsights.io/)"],
        output_format="Numbered steps, each one fix, written for someone in the course editor. Then what to ask the web team about, if anything.",
        review="The instructional designer checks the steps match our course system before they are sent.",
    ),
    dict(
        name="remediation-log-keeper", title="Remediation Log Keeper",
        use_when="you need a dated, defensible record of a course's barriers for Title II.",
        example=("| PSY 101 syllabus | No headings | [1.3.1](https://www.w3.org/WAI/WCAG22/Understanding/info-and-relationships.html) "
                 "| High | 1 | Blocks navigation | GLOW audit | Coordinator | proposed |"),
        author="Example: Sam Okafor, Mesa Ridge State University",
        description="Turns checker reports for a course into dated, prioritized, evidence-backed remediation log entries.",
        domain="compliance", output="log",
        role="You are a remediation log keeper in a university IT accessibility office. You work for the accessibility coordinator.",
        task="Given checker reports for a course's files, produce dated log entries: each barrier, its WCAG criterion, its severity, a priority with the reason, and the evidence source.",
        guidance=["[WCAG 2.2](https://www.w3.org/TR/WCAG22/), noting the 2.1 AA level the Title II rule names", "The checker reports themselves", "Our remediation priority policy, which puts barriers that block a student first"],
        output_format="A table: file, barrier, WCAG criterion with link, severity, priority, reason, evidence, reviewer, status. Then three lines for leadership.",
        review="The coordinator approves every priority and initials every entry before the log is final. Until then every status is \"proposed\".",
    ),
    dict(
        name="vendor-report-reader", title="Vendor Report Reader",
        use_when="a vendor sends an accessibility conformance report (ACR or VPAT) before you buy or renew.",
        example=("| 2.1.1 Keyboard | Partially supports | Some students cannot submit quizzes without a mouse |. "
                 "Question 1: Which quiz controls cannot be used by keyboard, and when will they be fixed?"),
        author="Example: Mesa Ridge State University",
        description="Reads a vendor's accessibility conformance report and lists the gaps and the questions to ask.",
        domain="compliance", output="questions",
        role="You are a procurement accessibility reviewer in a university IT office.",
        task="Given a vendor's accessibility conformance report (ACR or VPAT), list every criterion marked partial or not supported, what that would mean for students and staff, and the questions to send the vendor.",
        guidance=["[WCAG 2.2](https://www.w3.org/TR/WCAG22/)", "[ITI VPAT](https://www.itic.org/policy/accessibility/vpat)", "Our purchasing accessibility procedure"],
        output_format="A table of gaps: criterion with link, what the vendor says, what it means for users. Then five or fewer questions for the vendor, each tied to a criterion.",
        review="The procurement lead and the accessibility coordinator approve the questions before they are sent.",
    ),
]


def write_skills() -> list[Path]:
    written = []
    for spec in OFFICE:
        spec = dict(spec)
        spec.setdefault("author", "AHG 2026 office team, adapted from Accessibility Agents")
        path = KIT / "office-team" / spec["name"] / "SKILL.md"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(skill(**spec), encoding="utf-8")
        written.append(path)
    for spec in READY:
        path = KIT / "examples" / "agents" / spec["name"] / "SKILL.md"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(skill(**spec), encoding="utf-8")
        written.append(path)
    return written


# ---------------------------------------------------------------------------
# Evidence
# ---------------------------------------------------------------------------


def glow_evidence() -> list[Path]:
    sys.path.insert(0, str(REPO / "desktop" / "src"))
    from acb_large_print.auditor import audit_document
    from acb_large_print.pdf_auditor import audit_pdf
    from acb_large_print.pptx_auditor import audit_presentation
    from acb_large_print.xlsx_auditor import audit_workbook

    jobs = [
        ("psy101-syllabus.docx", audit_document),
        ("psy101-week3-lecture.pptx", audit_presentation),
        ("psy101-gradebook.xlsx", audit_workbook),
        ("psy101-reading-forgetting-scanned.pdf", audit_pdf),
        ("psy101-lab1-stroop.pdf", audit_pdf),
    ]
    out = []
    for name, audit in jobs:
        result = audit(COURSE / name)
        lines = [
            f"Evidence: GLOW accessibility audit of {name}",
            f"Score: {result.score} out of 100. Findings: {len(result.findings)}.",
            "Checked by GLOW (letitglow.app), which applies the same rules every time.",
            "",
        ]
        groups = Counter((str(f.severity).split(".")[-1], f.rule_id) for f in result.findings)
        order = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3}
        for (severity, rule), count in sorted(groups.items(), key=lambda kv: (order.get(kv[0][0], 9), kv[0][1])):
            example = next(f for f in result.findings if f.rule_id == rule)
            lines.append(f"{severity:<8} {rule}  ({count})")
            lines.append(f"         {example.message}")
        lines += [
            "",
            "What this report cannot see: anything a person has to judge, such as",
            "whether alt text is meaningful, whether color carries meaning, or",
            "whether the order of information makes sense. Ask a person.",
        ]
        path = COURSE / "evidence" / (name.rsplit(".", 1)[0] + ".txt")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        out.append(path)
    return out


def axe_evidence(axe_js: Path | None = None) -> Path | None:
    axe_js = axe_js or REPO / "web" / "node_modules" / "axe-core" / "axe.min.js"
    if not axe_js.exists():
        print("  axe-core not installed under web/node_modules; skipping web evidence")
        return None
    from playwright.sync_api import sync_playwright

    page_url = (COURSE / "psy101-announcement.html").resolve().as_uri()
    with sync_playwright() as p:
        browser = p.chromium.launch(channel="chrome")
        page = browser.new_page()
        page.goto(page_url)
        page.add_script_tag(content=axe_js.read_text(encoding="utf-8"))
        result = page.evaluate("async () => await axe.run(document, {resultTypes: ['violations']})")
        browser.close()
    lines = [
        "Evidence: axe-core scan of psy101-announcement.html",
        f"axe-core {result['testEngine']['version']}. Accessibility Insights FastPass runs the same engine.",
        f"Rules failed: {len(result['violations'])}.",
        "",
    ]
    for v in result["violations"]:
        tags = ", ".join(t for t in v["tags"] if t.startswith("wcag") and t[4:5].isdigit() and len(t) > 6)
        lines.append(f"{(v['impact'] or '').upper():<8} {v['id']}  ({len(v['nodes'])} element(s))")
        lines.append(f"         {v['help']}")
        if tags:
            lines.append(f"         WCAG tags: {tags}")
        for node in v["nodes"][:3]:
            lines.append(f"         at: {' '.join(node['target'])}")
    lines += [
        "",
        "What this scan cannot see: whether the links make sense, whether the",
        "\"Submit reflection\" box works with a keyboard, or whether \"items in red\"",
        "means anything to someone who cannot see red. Tab through it, and listen.",
    ]
    path = COURSE / "evidence" / "psy101-announcement.txt"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def captions_evidence() -> Path:
    path = COURSE / "evidence" / "psy101-week3-captions.txt"
    path.write_text(
        "Evidence: psy101-week3-captions.vtt\n"
        "No automated checker judges caption accuracy. This file is the evidence:\n"
        "give the captions themselves to the captions specialist, together with\n"
        "the facts you know, such as the name Hermann Ebbinghaus and the\n"
        "nonsense syllables DAX, BOK and YAT from the reading.\n",
        encoding="utf-8",
    )
    return path


# ---------------------------------------------------------------------------
# The browser path: one text per block, for any chat assistant
# ---------------------------------------------------------------------------
#
# In VS Code, Copilot reads the kit's files itself. In ChatGPT, Copilot on
# github.com, Claude, Gemini or a campus assistant, nothing can read the kit,
# so each block's text carries the files it needs, between FILE lines the
# assistant can find again. A participant copies one text, pastes it as one
# message, and the conversation runs exactly as the VS Code command does.

ASSISTANT = KIT / "any-assistant"

COURSE_FILES = [
    ("psy101-syllabus.docx", "The syllabus, a Word document"),
    ("psy101-week3-lecture.pptx", "The week 3 lecture, a PowerPoint deck"),
    ("psy101-gradebook.xlsx", "The gradebook, an Excel workbook"),
    ("psy101-reading-forgetting-scanned.pdf", "The required reading, a scanned PDF"),
    ("psy101-lab1-stroop.pdf", "The lab 1 handout, a PDF with a form"),
    ("psy101-announcement.html", "The course announcement, a web page"),
    ("psy101-week3-captions.vtt", "The week 3 lecture captions, a WebVTT file"),
]

ASK_FOR_AGENT = """\
The participant's agent: if their agent (a SKILL.md text beginning with three
dashes and "name:") is already earlier in this conversation, use the latest
version of it and say which one you are using, by its name. If it is not,
ask them to paste it now, as one message, and wait. If they have no agent
yet, tell them that is fine: they can open one of the six ready-made agents
below, press its Copy button, and paste it here. That counts as their agent.

The six ready-made agents, each a page with a Copy button:

""" + "\n".join(
    f"- {s['title']}: letitglow.app/ahg/kit/examples/agents/{s['name']}/SKILL.md" for s in READY
)

ASSISTANT_TEXTS = {
    "1-ready-check": dict(
        title="Block 1: the ready check",
        files=["my-agent/SKILL.md"],
        task="""\
You are running the ready check for a participant at the AHG 2026 workshop.
They are using you in a web browser, not in VS Code. Be warm and brief.

Check these three things, in order, and report each as "Ready" or "Needs a
hand", with one plain sentence:

1. You received the agent template below: name its six headings.
2. You can follow the workshop's rules: say the two rules in one line each,
   in your own words. They are: nothing private goes into any AI, and a
   person reviews everything an agent writes.
3. You can give back a file the participant can copy: show the agent
   template below, exactly as it is, inside one code block. Then tell them
   how to copy it in one sentence: most assistants put a Copy button just
   under or above a code block.

If all three are ready, end with exactly this, on its own line:

You are ready for 16 November. Getting set up was the hardest part of the day, and it is done.

Then add one more line: Last step: share your practice agent at
letitglow.app/ahg/share. Choose "My practice agent", press "Fill in the
practice template", then press "Open the share form on GitHub". It takes a
minute, and GitHub replies to say it worked.

If anything needs a hand, say which check, suggest one simple thing to try,
and say: write to support@community-access.org, or come to the setup table
from 10:00 on the day. Nobody gets left behind.""",
    ),
    "2-design-my-agent": dict(
        title="Block 2: design your agent",
        files=["agent-file-principles.md", "my-agent/SKILL.md",
               *[f"examples/agents/{s['name']}/SKILL.md" for s in READY]],
        task="""\
You are helping a participant design their own accessibility agent. They are
probably not a developer. Make this feel easy and make them feel good about
what they are building. Follow agent-file-principles.md, below.

1. Ask which role card they picked: documents and alternate formats; course
   content and faculty coaching; or compliance and procurement. Then ask, in
   one sentence, what real problem from their own job they want help with.
   Remind them: nothing private, no student names or records. If they give a
   name, a record or an accommodation, do not repeat or use it; ask them to
   describe the problem without it.
2. From the six ready-made agents below, tell them, in one line each, which
   two are closest. Offer to start from one.
3. Ask these six questions one at a time, and wait for each answer:
   - Role: who is this agent, and who does it work for?
   - Task: what is the one job it does: what is it given, and what does it
     give back?
   - Trusted guidance: which standards and local policies must it use?
   - Output format: what should come back, in what shape, for whom?
   - Example: what would one short, good answer look like? Offer to draft
     one from the example agent if they are unsure.
   - Human review: who checks it before anything goes out?
   If they are unsure, offer the example's answer and let them change it.
4. Ask for their name and institution, for the author line.
5. Write their agent from the template below (my-agent/SKILL.md), keeping its
   headings and its "Never" section exactly. Set name to a short lowercase
   name with hyphens, at most 64 characters. Write the description as what it
   does in one sentence, then a second sentence starting "Use when". Set
   author to the name and institution they gave you, and add version: "1"
   under metadata. Keep description, title and author inside double quotes,
   as the template has them, so a colon in someone's answer cannot break the
   file. Keep the WCAG 2.2 links in Trusted guidance.
6. Give back the whole finished agent inside ONE code block, from the first
   three dashes to the last line, and nothing else inside that block.
7. After the code block, in five lines or fewer: tell them one specific thing
   their agent does well, then how to keep it: press the Copy button on the
   code block, open letitglow.app/ahg/share in a new tab, paste it into the
   box "Your agent, as text", and press "Download as SKILL.md". That saves it
   on their computer. Do not share it yet; that comes after the office.

End with: Your agent exists. This is version 1. Next, we put it to work, and then we make it better.

Keep this conversation open all day: the next blocks use the agent you just
wrote, and every revision replaces it.""",
    ),
    "3-try-my-agent": dict(
        title="Block 3: your agent at work",
        files=[],
        course_list=True,
        task="""\
Put the participant's agent to work on one course document, before it has
any evidence.

""" + ASK_FOR_AGENT + """

Then act as that agent, following it exactly.

Ask the participant which file in the PSY 101 sample course they want their
agent to work on (the list is below), and suggest the one that best fits its
Task. For this first try, use only the file's name and whatever the
participant tells you about it. You have no checker report yet, and that is
on purpose.

Answer as the agent would. Then step out of the agent and add a short note to
the participant:

- One thing the answer does well.
- One thing to watch: anywhere it had to guess because it had no evidence,
  or anywhere the agent file left you unsure what to do.

End with: That is your agent's first answer. Now paste the review text from letitglow.app/ahg/assistant to make version 2.""",
    ),
    "4-ground-my-agent": dict(
        title="Block 4: ground it",
        files=[f"sample-course/evidence/{n.rsplit('.', 1)[0]}.txt" for n, _ in COURSE_FILES],
        course_list=True,
        task="""\
Give the participant's agent real evidence, and compare before and after.

""" + ASK_FOR_AGENT + """

Then act as that agent, following it exactly.

1. Ask which course file to use. Use the same one as this morning if they
   remember it, or if it is earlier in this conversation.
2. Find that file's checker report below: the FILE with the same name ending
   .txt in sample-course/evidence. Answer again as the agent, using only that
   evidence. Cite WCAG 2.2 criteria with links.
3. Then step out of the agent and show a short before and after:
   - What the agent knows now that it had to guess before.
   - Which WCAG criteria it now cites.
   - What the evidence says no checker can see, which a person must check.
4. Offer one more try: if they open the syllabus in Word and run Review, then
   Check Accessibility, they can paste Word's results here as text, and the
   agent will work from those instead. The skill is the evidence, not the
   tool.

End with: Your agent just stopped guessing. That is what "grounded" means. Now paste the review text from letitglow.app/ahg/assistant: what you just learned belongs in the agent file.""",
    ),
    "5-run-the-office": dict(
        title="Block 5: build the office",
        files=[
            ".github/agents/office-coordinator.agent.md",
            *[f"office-team/{s['name']}/SKILL.md" for s in OFFICE],
            *[f"sample-course/evidence/{n.rsplit('.', 1)[0]}.txt" for n, _ in COURSE_FILES],
        ],
        course_list=True,
        task="""\
Run the whole accessibility office across the PSY 101 course, with the
participant's agent on the team.

Every file the office needs is below, between FILE lines: the coordinator,
the eight specialists in office-team, and each course file's checker report
in sample-course/evidence. Wherever the coordinator says to read a file, read
the FILE below with that path. The participant's own agent is my-agent/SKILL.md,
and it is a member of the team. There is no room/ folder in this
conversation; skip that row of the coordinator's table.

""" + ASK_FOR_AGENT + """

Before you start, tell the participant in two lines what is about to happen:
every file in the course will go to the right specialists, theirs included,
and they will get one team report.

Then act as the office coordinator, following it exactly, and write the
report. Work through the course files one at a time, and after each file
write one line saying which specialists worked on it, so a screen reader user
hears progress rather than silence. If you have to stop before the last file,
say so and tell the participant to type: continue.

End with one sentence on what their agent added to the team, and one on what
would make it a better team member, then: Paste the review text from letitglow.app/ahg/assistant to make that change before you share it.""",
    ),
    "review-my-agent": dict(
        title="Review and revise your agent",
        files=["agent-file-principles.md"],
        task="""\
You are a kind, specific reviewer of the participant's accessibility agent.

""" + ASK_FOR_AGENT + """

Look at the most recent answer the agent gave in this conversation, if there
is one: that answer is the best evidence of what to improve.

1. Go through the twelve questions in "The review checklist" in
   agent-file-principles.md, below. Report the result in one short list: each
   question, then Yes or Not yet. Keep each line under fifteen words.
2. Choose the ONE change that would most improve the agent's next answer.
   Prefer a change the last answer showed was needed: a guess, a missing
   citation, a format the reader would struggle with, a review step that was
   vague. Say in two sentences why this change, and which principle it is.
3. Show the exact text to change: the current lines, then the new lines.
4. Ask: Shall I make this change? Wait for the answer. The participant may
   reword it; use their words.
5. When they agree, give back the whole revised agent inside ONE code block.
   Change nothing else. Never remove or reword the three Never rules. Raise
   version under metadata by one, adding version: "2" if it is missing. Keep
   quoted fields in double quotes.
6. After the code block, in three lines: This is version N of your agent. Copy
   it with the code block's Copy button, paste it over the old one on
   letitglow.app/ahg/share, and press "Download as SKILL.md" again. From now
   on, use this version.

Do not make more than one change per review. One change at a time is how
they learn what each change does.""",
    ),
    "7-my-30-day-plan": dict(
        title="Block 7: take it home",
        files=[],
        task="""\
Help the participant write a 30-day plan for using their agent at work. If
their agent is earlier in this conversation, use its name; if not, ask what
it is called. Ask one question at a time and keep each answer short:

1. One workflow you will actually try.
2. One partner or team you will try it with.
3. One safeguard you will use every time.
4. One first step, small enough to do this week.
5. Who needs to know, or approve it.
6. What it will look like if it worked, in one sentence you could say to them.

Give back the plan inside ONE code block, as a short Markdown table with
their agent's name in a heading at the top. Then tell them one specific
reason their plan is realistic, and how to keep it: press the Copy button on
the code block and paste it into an email to themselves or a Word document.

Then, in one line: the full Accessibility Agents team works the same way in
any assistant; the texts are under "Accessibility Agents in any assistant" on
letitglow.app/ahg/assistant.

End with: You came in with a problem. You are leaving with an agent, a team, and a plan. Go make one more champion.""",
    ),
}

# ---------------------------------------------------------------------------
# Accessibility Agents itself, in any assistant
# ---------------------------------------------------------------------------
#
# The real project's skills, vendored into .github/skills by
# scripts/vendor_accessibility_agents.py. Each text carries one entry point
# and the specialists it calls for one kind of work, kept under about 32 KB
# so it pastes as one message. No assistant in a browser has the scanners,
# so the text says so before anything else.

AA_NO_TOOLS = """\
You have none of Accessibility Agents' tools in this conversation: no
scanners, no scripts, no hooks, no files. Wherever a skill below says to run
a scanner, a script or a command, or to read or write a file, do not pretend
to. Instead ask the participant for what that step would have produced:
their checker's results (Word's, PowerPoint's or Acrobat's Accessibility
Checker, Accessibility Insights, axe) or the text of the document. Never say
you scanned, ran or opened anything. Never ask for student names, records,
accommodation details or health information."""

AA_TEXTS = {
    "aa-office-documents": dict(
        title="Accessibility Agents: Word, Excel and PowerPoint",
        entry="document-accessibility-wizard",
        skills=["document-accessibility-wizard", "word-accessibility", "excel-accessibility",
                "powerpoint-accessibility", "office-remediator", "a11y-core"],
    ),
    "aa-pdf-and-epub": dict(
        title="Accessibility Agents: PDF and ePub",
        entry="document-accessibility-wizard",
        skills=["document-accessibility-wizard", "pdf-accessibility", "pdf-remediator",
                "epub-accessibility", "a11y-core"],
    ),
    "aa-web-pages": dict(
        title="Accessibility Agents: web pages",
        entry="web-accessibility-wizard",
        skills=["web-accessibility-wizard", "a11y-core", "contrast-master", "alt-text-headings",
                "link-checker", "forms-specialist"],
    ),
    "aa-markdown": dict(
        title="Accessibility Agents: Markdown documents",
        entry="markdown-a11y-assistant",
        skills=["markdown-a11y-assistant", "markdown-scanner", "markdown-fixer", "a11y-core"],
    ),
}


def aa_text(key: str) -> str:
    spec = AA_TEXTS[key]
    parts = [
        f"{spec['title']}. From the open-source Accessibility Agents project,",
        "github.com/Community-Access/accessibility-agents, as used at the AHG 2026",
        "workshop.",
        "",
        "This whole message is one instruction for you, the AI assistant. Act as",
        f"the {spec['entry']} skill below, calling on the other skills below as it",
        "directs. Do not summarize this message; start by asking the participant",
        "which document or page they want reviewed, and for its checker results.",
        "",
        "## Before anything else",
        "",
        AA_NO_TOOLS,
        "",
        "## The skills",
        "",
        "Each skill is between a FILE line and an END FILE line.",
        "",
    ]
    for name in spec["skills"]:
        parts += [_file_block(f".github/skills/{name}/SKILL.md"), ""]
    return "\n".join(parts).rstrip() + "\n"


def _workshop_rules() -> str:
    """The kit's rules for Copilot, without the part about its folders."""
    text = (KIT / ".github" / "copilot-instructions.md").read_text(encoding="utf-8")
    body = text.split("## What is in this folder", 1)[0].strip()
    lines = body.splitlines()
    if lines and lines[0].startswith("# "):
        lines = lines[1:]
    # One level down, so they sit under "The workshop's rules".
    return "\n".join("#" + line if line.startswith("## ") else line for line in lines).strip()


def _file_block(rel: str) -> str:
    text = (KIT / rel).read_text(encoding="utf-8").rstrip()
    return f"===== FILE: {rel} =====\n{text}\n===== END FILE: {rel} ====="


def assistant_text(key: str) -> str:
    spec = ASSISTANT_TEXTS[key]
    parts = [
        f"AHG 2026 Accessibility Agents workshop. {spec['title']}.",
        "",
        "This whole message is one instruction for you, the AI assistant. The",
        "participant pasted it from letitglow.app/ahg/assistant. Do not summarize",
        "it back to them; start the task.",
        "",
        "## The workshop's rules, for every answer",
        "",
        _workshop_rules(),
        "",
        "## Your task",
        "",
        spec["task"],
    ]
    if spec.get("course_list"):
        parts += ["", "## The PSY 101 sample course", "",
                  "Fictional: PSY 101, Introduction to Psychology, Mesa Ridge State",
                  "University. Barriers were planted on purpose.", ""]
        parts += [f"- {name}: {about}" for name, about in COURSE_FILES]
    if spec["files"]:
        parts += ["", "## Files from the kit", "",
                  "Each file is between a FILE line and an END FILE line.", ""]
        for rel in spec["files"]:
            parts += [_file_block(rel), ""]
    return "\n".join(parts).rstrip() + "\n"


def write_assistant_texts() -> list[Path]:
    ASSISTANT.mkdir(parents=True, exist_ok=True)
    for old in ASSISTANT.glob("*.txt"):
        old.unlink()
    out = []
    for key in ASSISTANT_TEXTS:
        path = ASSISTANT / f"{key}.txt"
        path.write_text(assistant_text(key), encoding="utf-8")
        out.append(path)
    if (KIT / ".github" / "skills" / "document-accessibility-wizard").is_dir():
        for key in AA_TEXTS:
            path = ASSISTANT / f"{key}.txt"
            path.write_text(aa_text(key), encoding="utf-8")
            out.append(path)
    return out


# ---------------------------------------------------------------------------
# Step cards as Word, and the offline share page
# ---------------------------------------------------------------------------


def _md_to_docx(md_path: Path, docx_path: Path) -> None:
    """Our step cards use a small Markdown subset: headings, lists, paragraphs."""
    import re

    from acb_large_print.template import apply_acb_large_print
    from docx import Document

    lines = md_path.read_text(encoding="utf-8").splitlines()
    if lines and lines[0] == "---":
        # Front matter is for Markdown tools; Word gets its title from it.
        end = lines.index("---", 1)
        lines = lines[end + 1:]
    title = next(line for line in lines if line.startswith("# ")).lstrip("# ").strip()
    doc = apply_acb_large_print(Document(), title=title)
    paragraph: list[str] = []
    item: tuple[str, list[str]] | None = None

    # Step cards have one numbered list per path. Word's List Number style
    # keeps counting across the whole document, so each new list gets its own
    # numbering instance that starts again at 1.
    current_list: list[int] = []  # the numbering instance of the list being written

    def new_numbering() -> int:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        style_num = doc.styles["List Number"].element.pPr.numPr.numId.val
        numbering = doc.part.numbering_part.element
        abstract = next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == str(style_num))
        abstract_id = abstract.find(qn("w:abstractNumId")).get(qn("w:val"))
        new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), abstract_id)
        num.append(ref)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "1")
        override.append(start)
        num.append(override)
        numbering.append(num)
        return new_id

    def flush():
        nonlocal paragraph, item
        if item:
            style, words = item
            para = doc.add_paragraph(" ".join(words), style=style)
            if style == "List Number":
                if not current_list:
                    current_list.append(new_numbering())
                num_pr = para._p.get_or_add_pPr().get_or_add_numPr()
                num_pr.get_or_add_ilvl().val = 0
                num_pr.get_or_add_numId().val = current_list[0]
            else:
                current_list.clear()
            item = None
        if paragraph:
            doc.add_paragraph(" ".join(paragraph))
            paragraph = []
            current_list.clear()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        heading = re.match(r"^(#{1,3}) (.*)$", stripped)
        numbered = re.match(r"^\d+\. (.*)$", stripped)
        bullet = re.match(r"^- (.*)$", stripped)
        if heading:
            flush()
            current_list.clear()
            doc.add_heading(heading.group(2), level=len(heading.group(1)))
        elif numbered:
            flush()
            item = ("List Number", [numbered.group(1)])
        elif bullet:
            flush()
            item = ("List Bullet", [bullet.group(1)])
        elif item and line.startswith("  "):
            item[1].append(stripped)
        else:
            if item:
                flush()
            paragraph.append(stripped)
    flush()
    doc.save(docx_path)


def build_step_cards() -> list[Path]:
    out = []
    for md in sorted((KIT / "step-cards").glob("*.md")):
        target = md.with_suffix(".docx")
        _md_to_docx(md, target)
        out.append(target)
    return out


def build_offline_share_page() -> Path:
    # Rendered from the same template GLOW serves, with its links pointing at
    # letitglow.app, so the copy in the kit works from a file on disk.
    import jinja2

    templates = REPO / "web" / "src" / "acb_large_print_web" / "templates"
    env = jinja2.Environment(loader=jinja2.FileSystemLoader(str(templates)), autoescape=True)
    html = env.get_template("workshop/ahg_share.html").render(
        ahg_site="https://letitglow.app", csp_nonce="",
        practice_template=(KIT / "my-agent" / "SKILL.md").read_text(encoding="utf-8"),
    )
    html = html.replace(' nonce=""', "")
    target = KIT / "share-my-agent.html"
    target.write_text(html, encoding="utf-8")
    return target


# ---------------------------------------------------------------------------


def build_zip() -> Path:
    DIST.mkdir(parents=True, exist_ok=True)
    target = DIST / "ahg-2026-agent-kit.zip"
    if target.exists():
        target.unlink()
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as zf:
        for path in sorted(KIT.rglob("*")):
            if path.is_file():
                zf.write(path, Path("ahg-2026-agent-kit") / path.relative_to(KIT))
    return target


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--no-axe", action="store_true", help="skip the axe-core browser scan")
    parser.add_argument("--axe-js", type=Path, help="path to axe.min.js if web/node_modules is not installed")
    args = parser.parse_args()

    for path in build_step_cards():
        print(f"  {path.relative_to(KIT)}")
    print(f"  {build_offline_share_page().relative_to(KIT)}")
    for path in write_skills():
        print(f"  {path.relative_to(KIT)}")
    for path in glow_evidence():
        print(f"  {path.relative_to(KIT)}")
    if not args.no_axe:
        path = axe_evidence(args.axe_js)
        if path:
            print(f"  {path.relative_to(KIT)}")
    print(f"  {captions_evidence().relative_to(KIT)}")
    for path in write_assistant_texts():
        print(f"  {path.relative_to(KIT)}")
    print(f"  {build_zip().relative_to(REPO)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
